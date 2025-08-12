"""
Excel Export Module for Audit Reports
Generates comprehensive XLSX files with multiple tabs for journal entry testing
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from datetime import datetime
import logging
from typing import Dict, List, Any

logger = logging.getLogger(__name__)

class AuditExcelExporter:
    """Excel exporter for audit reports with multiple tabs"""
    
    def __init__(self):
        self.workbook = None
        
    def create_audit_report(self, data_file, analysis_results, output_path):
        """Create comprehensive audit report Excel file"""
        try:
            # DEBUG: Print data structure received by Excel exporter
            print("=" * 80)
            print("DEBUG: EXCEL EXPORTER RECEIVED DATA")
            print("=" * 80)
            print(f"Data File: {data_file.file_name}")
            print(f"Analysis Results Keys: {list(analysis_results.keys())}")
            
            for key, value in analysis_results.items():
                if isinstance(value, dict):
                    print(f"\n{key.upper()}:")
                    for sub_key, sub_value in value.items():
                        if isinstance(sub_value, list):
                            print(f"  {sub_key}: {len(sub_value)} items")
                            if sub_value and len(sub_value) > 0:
                                print(f"    First item sample: {sub_value[0]}")
                        elif isinstance(sub_value, dict):
                            print(f"  {sub_key}: {len(sub_value)} dict items")
                            if sub_value:
                                print(f"    Keys: {list(sub_value.keys())}")
                        else:
                            print(f"  {sub_key}: {sub_value}")
                else:
                    print(f"{key}: {type(value)}")
            
            print("=" * 80)
            print()
            
            # Create workbook
            self.workbook = Workbook()
            
            # Remove default sheet
            self.workbook.remove(self.workbook.active)
            
            # Create tabs
            self._create_instructions_tab()
            self._create_overview_tab(data_file, analysis_results)
            self._create_summary_tab(data_file, analysis_results)
            self._create_backdated_tab(analysis_results)
            self._create_closing_entries_tab(analysis_results)
            self._create_duplicate_tab(analysis_results)
            self._create_holiday_tab(analysis_results)
            self._create_user_analysis_tab(analysis_results)
            self._create_unusual_days_tab(analysis_results)
            self._create_detailed_audit_log_tab(data_file, analysis_results)
            
            # Save workbook
            self.workbook.save(output_path)
            logger.info(f"Audit report saved to: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Error creating audit report: {e}")
            raise
    
    def _create_instructions_tab(self):
        """Create Instructions tab with audit methodology"""
        ws = self.workbook.create_sheet("Instructions")
        
        # Title
        ws['A1'] = "JOURNAL ENTRY TESTING — JOURNAL ENTRY TESTING TEMPLATE — COMBINATION APPROACH"
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:Z1')
        
        # Header styling
        header_fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
        ws['A1'].fill = header_fill
        ws['A1'].font = Font(color="FFFFFF", bold=True)
        
        # Instructions content
        instructions = [
            ("HOW TO USE THIS WORKBOOK", ""),
            ("Introduction", "This workbook is designed to assist engagement teams in documenting journal entry testing procedures using the combination approach."),
            ("Note", "This template is only to be used by engagement teams using the combination approach for journal entry testing."),
            ("Steps within this workbook:", ""),
            ("", ""),
            ("Tab 1. Understanding the Financial Reporting Process (FRP) and Controls Over Journal Entries", ""),
            ("Step 1", "Obtain an Understanding of the FRP and Controls Over Journal Entries"),
            ("Step 2", "The Effectiveness of Controls Over Journal Entries and Its Impact on the Audit"),
            ("", ""),
            ("Tab 2. Inquiries of FRP Personnel", ""),
            ("Step 3", "Inquiries of Entity Personnel Involved in the FRP"),
            ("", ""),
            ("Tab 3. Top-Down Journal Entry Testing Scoping Considerations", ""),
            ("Step 4", "Auditor's Assessment of Fraud Risks"),
            ("Step 5", "Incorporating Unpredictability"),
            ("Step 6", "Scoping of Components"),
            ("", ""),
            ("Tab 4. Design of Journal Entry Testing Approach", ""),
            ("Step 7", "Excluding Accounts from Journal Entry Testing"),
            ("Step 8", "Determine Whether to Test Journal Entries Recorded Throughout the Period"),
            ("Step 9", "Summary of Audit Documentation When Using the Combination Approach"),
            ("", ""),
            ("Tab 5. Documentation of Judgments Related to Journal Entry Tests", ""),
            ("Step 10", "Judgments Supporting Combination Approach Tests"),
            ("", ""),
            ("Tab 6. Testing Journal Entries That Exhibit Characteristics of Fraudulent Entries", ""),
            ("Step 11", "Testing Journal Entries That Exhibit Characteristics of Fraudulent Entries"),
            ("", ""),
            ("Note:", "In scenarios where activities within the journal entry testing process are performed by the entity, additional procedures may be required.")
        ]
        
        for i, (step, description) in enumerate(instructions, start=3):
            ws[f'A{i}'] = step
            ws[f'B{i}'] = description
            
            if step and not step.startswith("Step") and step != "Note:":
                ws[f'A{i}'].font = Font(bold=True)
                if step == "HOW TO USE THIS WORKBOOK":
                    ws[f'A{i}'].fill = header_fill
                    ws[f'A{i}'].font = Font(color="FFFFFF", bold=True)
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 80
    
    def _create_overview_tab(self, data_file, analysis_results):
        """Create Overview tab with summary statistics"""
        ws = self.workbook.create_sheet("Overview")
        
        # Title
        ws['A1'] = "AUDIT OVERVIEW"
        ws['A1'].font = Font(bold=True, size=16)
        ws.merge_cells('A1:D1')
        
        # Client Information
        ws['A3'] = "Client Name:"
        ws['B3'] = data_file.client_name
        ws['A4'] = "Company:"
        ws['B4'] = data_file.company_name
        ws['A5'] = "Engagement ID:"
        ws['B5'] = getattr(data_file, 'engagement_id', 'ENG-007')  # Use default if not available
        ws['A6'] = "Fiscal Year:"
        ws['B6'] = getattr(data_file, 'fiscal_year', '2025')  # Use default if not available
        ws['A7'] = "Audit Period:"
        ws['B7'] = f"{getattr(data_file, 'audit_start_date', '2025-01-01')} to {getattr(data_file, 'audit_end_date', '2025-12-31')}"
        
        # File Statistics
        ws['A9'] = "FILE STATISTICS"
        ws['A9'].font = Font(bold=True, size=14)
        ws['A10'] = "Total Records:"
        ws['B10'] = data_file.total_records
        ws['A11'] = "Processed Records:"
        ws['B11'] = getattr(data_file, 'processed_records', data_file.total_records)
        ws['A12'] = "Failed Records:"
        ws['B12'] = getattr(data_file, 'failed_records', 0)
        
        # Analysis Summary
        ws['A14'] = "ANALYSIS SUMMARY"
        ws['A14'].font = Font(bold=True, size=14)
        
        # Get analysis counts from real data
        total_anomalies = 0
        anomaly_breakdown = {}
        
        if 'duplicate_analysis' in analysis_results:
            duplicate_count = len(analysis_results['duplicate_analysis'].get('duplicate_list', []))
            if duplicate_count > 0:
                anomaly_breakdown['Duplicate Entries'] = duplicate_count
                total_anomalies += duplicate_count
            
        if 'backdated_analysis' in analysis_results:
            backdated_count = len(analysis_results['backdated_analysis'].get('backdated_entries', []))
            if backdated_count > 0:
                anomaly_breakdown['Backdated Entries'] = backdated_count
                total_anomalies += backdated_count
            
        if 'closing_entries_analysis' in analysis_results:
            closing_count = len(analysis_results['closing_entries_analysis'].get('closing_entries', []))
            if closing_count > 0:
                anomaly_breakdown['Closing Entries'] = closing_count
                total_anomalies += closing_count
            
        if 'holiday_analysis' in analysis_results:
            holiday_count = len(analysis_results['holiday_analysis'].get('holiday_postings', []))
            if holiday_count > 0:
                anomaly_breakdown['Holiday Postings'] = holiday_count
                total_anomalies += holiday_count
        
        if 'unusual_days_analysis' in analysis_results:
            unusual_count = len(analysis_results['unusual_days_analysis'].get('unusual_days', []))
            if unusual_count > 0:
                anomaly_breakdown['Unusual Days'] = unusual_count
                total_anomalies += unusual_count
        
        # Add user anomalies to overview
        if 'user_analysis' in analysis_results:
            user_anomalies = analysis_results['user_analysis'].get('user_anomalies', [])
            user_anomaly_count = len(user_anomalies)
            if user_anomaly_count > 0:
                anomaly_breakdown['User Anomalies'] = user_anomaly_count
                total_anomalies += user_anomaly_count
        
        ws['A15'] = "Total Anomalies Detected:"
        ws['B15'] = total_anomalies
        
        row = 16
        for anomaly_type, count in anomaly_breakdown.items():
            ws[f'A{row}'] = f"{anomaly_type}:"
            ws[f'B{row}'] = count
            row += 1
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 30
    
    def _create_summary_tab(self, data_file, analysis_results):
        """Create Summary tab with exception journal summary"""
        ws = self.workbook.create_sheet("Summary")
        
        # Title
        ws['A1'] = "Summary of Entries Flagged"
        ws['A1'].font = Font(bold=True, size=16)
        ws.merge_cells('A1:C1')
        
        # Client info
        ws['A2'] = "Client Name:"
        ws['B2'] = data_file.client_name
        ws['A3'] = "Testing Period:"
        ws['B3'] = getattr(data_file, 'fiscal_year', '2025')
        
        # Description
        ws['A5'] = "This worksheet shows a summary of the entries submitted to testing, using the Journal Entry Testing methodology."
        ws.merge_cells('A5:C5')
        
        # Exception Journal Summary Table
        ws['A7'] = "Exception Journal Summary"
        ws['A7'].font = Font(bold=True, size=14)
        ws.merge_cells('A7:C7')
        
        # Table headers
        headers = ["Test Type", "Test Performed", "Entries Samples"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=9, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # Test results using real data
        test_results = [
            ("Backdated Entries", "Yes", self._get_anomaly_count(analysis_results, 'backdated_analysis', 'backdated_entries')),
            ("Closing Entries", "Yes", self._get_anomaly_count(analysis_results, 'closing_entries_analysis', 'closing_entries')),
            ("Potential Duplicates", "Yes", self._get_anomaly_count(analysis_results, 'duplicate_analysis', 'duplicate_list')),
            ("Journals Posted on Holidays", "Yes", self._get_anomaly_count(analysis_results, 'holiday_analysis', 'holiday_postings')),
            ("Journals Posted to Seldom Used Accounts", "Yes", self._get_anomaly_count(analysis_results, 'user_analysis', 'seldom_used_accounts')),
            ("Unexpected AP Postings", "No", "-"),
            ("Unusual Account Activity", "No", "-"),
            ("Journals with Unusual Amounts", "Yes", self._get_anomaly_count(analysis_results, 'overall_analysis', 'unusual_amounts')),
            ("Journals Posted on Unusual Days", "Yes", self._get_anomaly_count(analysis_results, 'unusual_days_analysis', 'unusual_days')),
            ("Unusual Postings to Revenue", "No", "-"),
            ("Journals with Unusual Descriptions", "No", "-"),
            ("Journals with Unusual Offset", "No", "-"),
            ("User Analysis", "Yes", self._get_anomaly_count(analysis_results, 'user_analysis', 'user_transaction_summary')),
            ("Journals Containing Words of Interest", "Yes", self._get_anomaly_count(analysis_results, 'overall_analysis', 'words_of_interest')),
            ("Journals Containing Debits to Revenue", "No", "-")
        ]
        
        row = 10
        total_samples = 0
        for test_type, performed, count in test_results:
            ws.cell(row=row, column=1, value=test_type)
            ws.cell(row=row, column=2, value=performed)
            ws.cell(row=row, column=3, value=count)
            
            if count != "-" and count > 0:
                total_samples += count
            row += 1
        
        # Total row
        ws.cell(row=row, column=1, value="Total")
        ws.cell(row=row, column=1).font = Font(bold=True)
        ws.cell(row=row, column=3, value=total_samples)
        ws.cell(row=row, column=3).font = Font(bold=True)
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 40
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
    
    def _create_anomaly_tab(self, ws, data_list, tab_name, test_name):
        """Generic method to create anomaly tabs with core SAPGLPosting fields plus audit process fields"""
        
        # Title
        ws['A1'] = f"{tab_name.upper()} - JOURNAL ENTRY TESTING"
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:Z1')
        
        # Header styling
        header_fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
        ws['A1'].fill = header_fill
        ws['A1'].font = Font(color="FFFFFF", bold=True)
        
        # Description
        ws['A2'] = f"Test Name: {test_name}"
        ws['A2'].font = Font(bold=True)
        ws.merge_cells('A2:Z2')
        
        # Headers with core SAPGLPosting fields plus audit process fields
        headers = [
            # Core SAPGLPosting Fields
            "Document Number", "Posting Date", "GL Account", "Amount Local Currency", 
            "Transaction Type", "Local Currency", "Text", "Document Date", 
            "Offsetting Account", "User Name", "Entry Date", "Document Type",
            "Profit Center", "Cost Center", "Clearing Document", "Segment", 
            "WBS Element", "Plant", "Material", "Invoice Reference", "Billing Document",
            "Sales Document", "Purchasing Document", "Order Number", "Asset Number",
            "Network", "Assignment", "Tax Code", "Fiscal Year", "Posting Period", "Year Month",
            
            # Audit Process Fields
            "Test Name", "Voucher Prepared (Accountant)", "Authorized and Reviewed (Accounts)",
            "Entry in GL", "GL Update Accurately", "Evaluation of Audit Evidence Given the Fraud Characteristic(s) Identified", "Conclusion"
        ]
        
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Write data
        row = 4
        for entry in data_list:
            # Extract transaction data based on the structure
            transaction_data = self._extract_transaction_data(entry, tab_name)
            
            # Map core SAPGLPosting fields to columns
            col = 1
            ws.cell(row=row, column=col, value=transaction_data.get('document_number', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('posting_date', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('gl_account', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('amount_local_currency', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('transaction_type', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('local_currency', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('text', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('document_date', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('offsetting_account', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('user_name', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('entry_date', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('document_type', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('profit_center', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('cost_center', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('clearing_document', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('segment', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('wbs_element', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('plant', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('material', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('invoice_reference', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('billing_document', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('sales_document', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('purchasing_document', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('order_number', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('asset_number', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('network', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('assignment', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('tax_code', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('fiscal_year', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('posting_period', '')); col += 1
            ws.cell(row=row, column=col, value=transaction_data.get('year_month', '')); col += 1
            
            # Audit Process Fields
            ws.cell(row=row, column=col, value=test_name); col += 1  # Test Name
            ws.cell(row=row, column=col, value=''); col += 1  # Voucher Prepared (Accountant) - for auditor to fill
            ws.cell(row=row, column=col, value=''); col += 1  # Authorized and Reviewed (Accounts) - for auditor to fill
            ws.cell(row=row, column=col, value=''); col += 1  # Entry in GL - for auditor to fill
            ws.cell(row=row, column=col, value=''); col += 1  # GL Update Accurately - for auditor to fill
            ws.cell(row=row, column=col, value=''); col += 1  # Evaluation of Audit Evidence - for auditor to fill
            ws.cell(row=row, column=col, value=''); col += 1  # Conclusion - for auditor to fill
            
            row += 1
        
        # Add summary row
        summary_row = row
        ws.cell(row=summary_row, column=1, value=f"Total {tab_name}: {len(data_list)}")
        ws.cell(row=summary_row, column=1).font = Font(bold=True)
        ws.merge_cells(f'A{summary_row}:Z{summary_row}')
        
        # Adjust column widths for better readability
        column_widths = [
            15, 12, 15, 15, 12, 10, 30, 12, 15, 20, 12, 10, 15, 15, 15, 10, 15, 10, 15, 15, 15, 15, 15, 15, 15, 15, 10, 10, 15, 10, 10,
            20, 25, 25, 15, 20, 50, 15
        ]
        for col, width in enumerate(column_widths, 1):
            col_letter = get_column_letter(col)
            ws.column_dimensions[col_letter].width = width
    
    def _extract_transaction_data(self, entry, tab_name):
        """Extract transaction data from various nested structures and fetch complete SAPGLPosting data"""
        from core.models import SAPGLPosting
        
        transaction_id = None
        document_number = None
        
        if tab_name == "Duplicate Entries":
            if 'transaction1' in entry:
                transaction_data = entry['transaction1']
                transaction_id = transaction_data.get('id')
                document_number = transaction_data.get('document_number')
            elif 'transaction2' in entry:
                transaction_data = entry['transaction2']
                transaction_id = transaction_data.get('id')
                document_number = transaction_data.get('document_number')
        else:
            transaction_id = entry.get('transaction_id') or entry.get('id')
            document_number = entry.get('document_number')
        
        sap_posting = None
        if transaction_id:
            try:
                sap_posting = SAPGLPosting.objects.get(id=transaction_id)
            except SAPGLPosting.DoesNotExist:
                pass
        
        if not sap_posting and document_number:
            try:
                sap_posting = SAPGLPosting.objects.filter(document_number=document_number).first()
            except:
                pass
        
        if sap_posting:
            # Handle empty document_number by using alternative fields or generating one
            document_number = sap_posting.document_number
            if not document_number or document_number == '':
                if sap_posting.clearing_document:
                    document_number = f"CLR-{sap_posting.clearing_document}"
                elif sap_posting.invoice_reference:
                    document_number = f"INV-{sap_posting.invoice_reference}"
                elif sap_posting.billing_document:
                    document_number = f"BILL-{sap_posting.billing_document}"
                elif sap_posting.sales_document:
                    document_number = f"SALES-{sap_posting.sales_document}"
                elif sap_posting.purchasing_document:
                    document_number = f"PUR-{sap_posting.purchasing_document}"
                elif sap_posting.order_number:
                    document_number = f"ORD-{sap_posting.order_number}"
                elif sap_posting.asset_number:
                    document_number = f"ASSET-{sap_posting.asset_number}"
                elif sap_posting.network:
                    document_number = f"NET-{sap_posting.network}"
                else:
                    document_number = f"DOC-{str(sap_posting.id)[:8]}-{sap_posting.posting_date.strftime('%Y%m%d') if sap_posting.posting_date else 'NODATE'}"
            
            return {
                'document_number': document_number,
                'posting_date': sap_posting.posting_date,
                'gl_account': sap_posting.gl_account,
                'amount_local_currency': sap_posting.amount_local_currency,
                'transaction_type': sap_posting.transaction_type,
                'local_currency': sap_posting.local_currency,
                'text': sap_posting.text,
                'document_date': sap_posting.document_date,
                'offsetting_account': sap_posting.offsetting_account,
                'user_name': sap_posting.user_name,
                'entry_date': sap_posting.entry_date,
                'document_type': sap_posting.document_type,
                'profit_center': sap_posting.profit_center,
                'cost_center': sap_posting.cost_center,
                'clearing_document': sap_posting.clearing_document,
                'segment': sap_posting.segment,
                'wbs_element': sap_posting.wbs_element,
                'plant': sap_posting.plant,
                'material': sap_posting.material,
                'invoice_reference': sap_posting.invoice_reference,
                'billing_document': sap_posting.billing_document,
                'sales_document': sap_posting.sales_document,
                'purchasing_document': sap_posting.purchasing_document,
                'order_number': sap_posting.order_number,
                'asset_number': sap_posting.asset_number,
                'network': sap_posting.network,
                'assignment': sap_posting.assignment,
                'tax_code': sap_posting.tax_code,
                'fiscal_year': sap_posting.fiscal_year,
                'posting_period': sap_posting.posting_period,
                'year_month': sap_posting.year_month,
            }
        
        # If we couldn't find the complete record, return the partial data from the analysis
        if tab_name == "Duplicate Entries":
            if 'transaction1' in entry:
                return entry['transaction1']
            elif 'transaction2' in entry:
                return entry['transaction2']
        elif tab_name == "Backdated Entries":
            return entry
        elif tab_name == "Closing Entries":
            return entry
        elif tab_name == "Holiday Entries":
            return entry
        elif tab_name == "Unusual Days":
            return entry
        elif tab_name == "User Analysis":
            # For user analysis, we need to find actual SAPGLPosting records for this user
            # The entry contains user summary data, not individual transaction data
            user_name = entry.get('user', '')
            if user_name:
                # Find a sample transaction for this user to get the SAPGLPosting structure
                try:
                    sample_posting = SAPGLPosting.objects.filter(user_name=user_name).first()
                    if sample_posting:
                        # Generate document number
                        document_number = sample_posting.document_number
                        if not document_number or document_number == '':
                            if sample_posting.clearing_document:
                                document_number = f"CLR-{sample_posting.clearing_document}"
                            elif sample_posting.invoice_reference:
                                document_number = f"INV-{sample_posting.invoice_reference}"
                            elif sample_posting.billing_document:
                                document_number = f"BILL-{sample_posting.billing_document}"
                            elif sample_posting.sales_document:
                                document_number = f"SALES-{sample_posting.sales_document}"
                            elif sample_posting.purchasing_document:
                                document_number = f"PUR-{sample_posting.purchasing_document}"
                            elif sample_posting.order_number:
                                document_number = f"ORD-{sample_posting.order_number}"
                            elif sample_posting.asset_number:
                                document_number = f"ASSET-{sample_posting.asset_number}"
                            elif sample_posting.network:
                                document_number = f"NET-{sample_posting.network}"
                            else:
                                document_number = f"DOC-{str(sample_posting.id)[:8]}-{sample_posting.posting_date.strftime('%Y%m%d') if sample_posting.posting_date else 'NODATE'}"
                        
                        return {
                            'document_number': document_number,
                            'posting_date': sample_posting.posting_date,
                            'gl_account': sample_posting.gl_account,
                            'amount_local_currency': sample_posting.amount_local_currency,
                            'transaction_type': sample_posting.transaction_type,
                            'local_currency': sample_posting.local_currency,
                            'text': sample_posting.text,
                            'document_date': sample_posting.document_date,
                            'offsetting_account': sample_posting.offsetting_account,
                            'user_name': sample_posting.user_name,
                            'entry_date': sample_posting.entry_date,
                            'document_type': sample_posting.document_type,
                            'profit_center': sample_posting.profit_center,
                            'cost_center': sample_posting.cost_center,
                            'clearing_document': sample_posting.clearing_document,
                            'segment': sample_posting.segment,
                            'wbs_element': sample_posting.wbs_element,
                            'plant': sample_posting.plant,
                            'material': sample_posting.material,
                            'invoice_reference': sample_posting.invoice_reference,
                            'billing_document': sample_posting.billing_document,
                            'sales_document': sample_posting.sales_document,
                            'purchasing_document': sample_posting.purchasing_document,
                            'order_number': sample_posting.order_number,
                            'asset_number': sample_posting.asset_number,
                            'network': sample_posting.network,
                            'assignment': sample_posting.assignment,
                            'tax_code': sample_posting.tax_code,
                            'fiscal_year': sample_posting.fiscal_year,
                            'posting_period': sample_posting.posting_period,
                            'year_month': sample_posting.year_month,
                        }
                except:
                    pass
            
            # If we can't find a sample transaction, return a minimal structure
            return {
                'document_number': f"USER-{user_name}",
                'posting_date': None,
                'gl_account': entry.get('details', ''),
                'amount_local_currency': 0,
                'transaction_type': entry.get('anomaly_type', ''),
                'local_currency': '',
                'text': entry.get('details', ''),
                'document_date': None,
                'offsetting_account': '',
                'user_name': user_name,
                'entry_date': None,
                'document_type': '',
                'profit_center': '',
                'cost_center': '',
                'clearing_document': '',
                'segment': '',
                'wbs_element': '',
                'plant': '',
                'material': '',
                'invoice_reference': '',
                'billing_document': '',
                'sales_document': '',
                'purchasing_document': '',
                'order_number': '',
                'asset_number': '',
                'network': '',
                'assignment': '',
                'tax_code': '',
                'fiscal_year': '',
                'posting_period': '',
                'year_month': '',
            }
        
        # Default: return the entry as is
        return entry
    
    def _create_backdated_tab(self, analysis_results):
        """Create Backdated Entries tab"""
        backdated_data = analysis_results.get('backdated_analysis', {}).get('backdated_entries', [])
        
        if backdated_data:
            ws = self.workbook.create_sheet("Backdated Entries")
            self._create_anomaly_tab(ws, backdated_data, "Backdated Entries", "Backdated Entry Detection")
        else:
            ws = self.workbook.create_sheet("Backdated Entries")
            ws['A1'] = "No backdated entries found"
    
    def _create_closing_entries_tab(self, analysis_results):
        """Create Closing Entries tab"""
        closing_data = analysis_results.get('closing_entries_analysis', {}).get('closing_entries', [])
        
        if closing_data:
            ws = self.workbook.create_sheet("Closing Entries")
            self._create_anomaly_tab(ws, closing_data, "Closing Entries", "Closing Entry Detection")
        else:
            ws = self.workbook.create_sheet("Closing Entries")
            ws['A1'] = "No closing entries found"
    
    def _create_duplicate_tab(self, analysis_results):
        """Create Duplicate Entries tab"""
        duplicate_data = analysis_results.get('duplicate_analysis', {}).get('duplicate_list', [])
        
        if duplicate_data:
            ws = self.workbook.create_sheet("Duplicate Entries")
            self._create_anomaly_tab(ws, duplicate_data, "Duplicate Entries", "Duplicate Entry Detection")
        else:
            ws = self.workbook.create_sheet("Duplicate Entries")
            ws['A1'] = "No duplicate entries found"
    
    def _create_holiday_tab(self, analysis_results):
        """Create Holiday Entries tab"""
        holiday_data = analysis_results.get('holiday_analysis', {}).get('holiday_postings', [])
        
        if holiday_data:
            ws = self.workbook.create_sheet("Holiday Entries")
            self._create_anomaly_tab(ws, holiday_data, "Holiday Entries", "Holiday Entry Detection")
        else:
            ws = self.workbook.create_sheet("Holiday Entries")
            ws['A1'] = "No holiday entries found"
    
    def _create_unusual_days_tab(self, analysis_results):
        """Create Unusual Days tab"""
        unusual_data = analysis_results.get('unusual_days_analysis', {}).get('unusual_days', [])
        
        if unusual_data:
            ws = self.workbook.create_sheet("Unusual Days")
            self._create_anomaly_tab(ws, unusual_data, "Unusual Days", "Unusual Day Detection")
        else:
            ws = self.workbook.create_sheet("Unusual Days")
            ws['A1'] = "No unusual days entries found"
    
    def _create_user_analysis_tab(self, analysis_results):
        """Create User Analysis tab"""
        user_anomalies_data = analysis_results.get('user_analysis', {}).get('user_anomalies', [])
        
        if user_anomalies_data:
            ws = self.workbook.create_sheet("User Analysis")
            self._create_anomaly_tab(ws, user_anomalies_data, "User Analysis", "User Analysis Detection")
        else:
            ws = self.workbook.create_sheet("User Analysis")
            ws['A1'] = "No user anomalies found"
    
    def _create_detailed_audit_log_tab(self, data_file, analysis_results):
        """Create Detailed Audit Log tab with SAPGLPosting fields plus audit process fields"""
        ws = self.workbook.create_sheet("Detailed Audit Log")
        
        # Title
        ws['A1'] = "Detailed Audit Log - Working Paper"
        ws['A1'].font = Font(bold=True, size=16)
        ws.merge_cells('A1:Z1')
        
        # Description
        ws['A3'] = "This worksheet contains all flagged transactions for detailed audit review and testing."
        ws.merge_cells('A3:Z3')
        
        # Headers with core SAPGLPosting fields plus audit process fields
        headers = [
            # Core SAPGLPosting Fields
            "Document Number", "Posting Date", "GL Account", "Amount Local Currency", 
            "Transaction Type", "Local Currency", "Text", "Document Date", 
            "Offsetting Account", "User Name", "Entry Date", "Document Type",
            "Profit Center", "Cost Center", "Clearing Document", "Segment", 
            "WBS Element", "Plant", "Material", "Invoice Reference", "Billing Document",
            "Sales Document", "Purchasing Document", "Order Number", "Asset Number",
            "Network", "Assignment", "Tax Code", "Fiscal Year", "Posting Period", "Year Month",
            
            # Audit Process Fields
            "Test Name", "Voucher Prepared (Accountant)", "Authorized and Reviewed (Accounts)",
            "Entry in GL", "GL Update Accurately", "Evaluation of Audit Evidence Given the Fraud Characteristic(s) Identified", "Conclusion"
        ]
        
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Collect all entries from different analysis types
        all_entries = []
        
        # Add backdated entries
        if 'backdated_analysis' in analysis_results:
            backdated_entries = analysis_results['backdated_analysis'].get('backdated_entries', [])
            for entry in backdated_entries:
                entry_copy = entry.copy()
                entry_copy['test_name'] = 'Backdated Entries'
                all_entries.append(entry_copy)
        
        # Add closing entries
        if 'closing_entries_analysis' in analysis_results:
            closing_entries = analysis_results['closing_entries_analysis'].get('closing_entries', [])
            for entry in closing_entries:
                entry_copy = entry.copy()
                entry_copy['test_name'] = 'Closing Entries'
                all_entries.append(entry_copy)
        
        # Add holiday entries
        if 'holiday_analysis' in analysis_results:
            holiday_entries = analysis_results['holiday_analysis'].get('holiday_postings', [])
            for entry in holiday_entries:
                entry_copy = entry.copy()
                entry_copy['test_name'] = 'Holiday Entries'
                all_entries.append(entry_copy)
        
        # Add unusual days entries
        if 'unusual_days_analysis' in analysis_results:
            unusual_entries = analysis_results['unusual_days_analysis'].get('unusual_days', [])
            for entry in unusual_entries:
                entry_copy = entry.copy()
                entry_copy['test_name'] = 'Unusual Days'
                all_entries.append(entry_copy)
        
        # Add duplicate entries
        if 'duplicate_analysis' in analysis_results:
            duplicate_entries = analysis_results['duplicate_analysis'].get('duplicate_list', [])
            for entry in duplicate_entries:
                entry_copy = entry.copy()
                entry_copy['test_name'] = 'Duplicate Entries'
                all_entries.append(entry_copy)
        
        # Add user anomalies
        if 'user_analysis' in analysis_results:
            user_anomalies = analysis_results['user_analysis'].get('user_anomalies', [])
            for entry in user_anomalies:
                entry_copy = entry.copy()
                entry_copy['test_name'] = 'User Anomalies'
                all_entries.append(entry_copy)
        
        # Limit to first 100 entries for performance
        all_entries = all_entries[:100]
        
        # Write data
        row = 6
        for entry in all_entries:
            # Map core SAPGLPosting fields to columns
            col = 1
            ws.cell(row=row, column=col, value=entry.get('document_number', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('posting_date', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('gl_account', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('amount_local_currency', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('transaction_type', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('local_currency', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('text', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('document_date', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('offsetting_account', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('user_name', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('entry_date', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('document_type', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('profit_center', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('cost_center', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('clearing_document', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('segment', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('wbs_element', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('plant', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('material', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('invoice_reference', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('billing_document', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('sales_document', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('purchasing_document', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('order_number', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('asset_number', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('network', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('assignment', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('tax_code', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('fiscal_year', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('posting_period', '')); col += 1
            ws.cell(row=row, column=col, value=entry.get('year_month', '')); col += 1
            
            # Audit Process Fields
            ws.cell(row=row, column=col, value=entry.get('test_name', '')); col += 1  # Test Name
            ws.cell(row=row, column=col, value=''); col += 1  # Voucher Prepared (Accountant) - for auditor to fill
            ws.cell(row=row, column=col, value=''); col += 1  # Authorized and Reviewed (Accounts) - for auditor to fill
            ws.cell(row=row, column=col, value=''); col += 1  # Entry in GL - for auditor to fill
            ws.cell(row=row, column=col, value=''); col += 1  # GL Update Accurately - for auditor to fill
            ws.cell(row=row, column=col, value=''); col += 1  # Evaluation of Audit Evidence Given the Fraud Characteristic(s) Identified - for auditor to fill
            ws.cell(row=row, column=col, value=''); col += 1  # Conclusion - for auditor to fill
            
            row += 1
        
        # Add summary
        summary_row = row + 2
        ws.cell(row=summary_row, column=1, value="Summary:")
        ws.cell(row=summary_row, column=1).font = Font(bold=True)
        ws.cell(row=summary_row, column=2, value=f"Total Entries in Log: {len(all_entries)}")
        ws.cell(row=summary_row, column=2).font = Font(bold=True)
        ws.cell(row=summary_row, column=3, value="(Limited to first 100 entries for performance)")
        ws.cell(row=summary_row, column=3).font = Font(italic=True)
        
        # Adjust column widths for better readability
        column_widths = [
            15, 12, 15, 15, 12, 10, 30, 12, 15, 20, 12, 10, 15, 15, 15, 10, 15, 10, 15, 15, 15, 15, 15, 15, 15, 15, 10, 10, 15, 10, 10,
            20, 25, 25, 15, 20, 50, 15
        ]
        for col, width in enumerate(column_widths, 1):
            col_letter = get_column_letter(col)
            ws.column_dimensions[col_letter].width = width
    
    def _get_anomaly_count(self, analysis_results, analysis_type, field_name):
        """Helper method to get anomaly count from analysis results"""
        if analysis_type in analysis_results:
            data = analysis_results[analysis_type]
            if field_name in data:
                count = len(data[field_name])
                return count if count > 0 else "-"
        return "-"
